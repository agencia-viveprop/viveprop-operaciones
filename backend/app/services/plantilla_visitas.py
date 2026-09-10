"""La estructura del archivo de visitas, y una plantilla vacía con esos encabezados.

**El archivo no se llena a mano: sale de la consola de administración.** Por eso la
plantilla no sirve para tipear visitas, sino para **comparar** encabezados cuando la
carga falla y no se entiende por qué --el mismo motivo que tiene la de canjes--.

Los nombres viven en `importar_visitas.COLUMNAS_REQUERIDAS`, que es lo que la carga
verifica de verdad. Acá se les agrega el grupo y la ayuda para la pantalla.

**Solo `Propiedad` es obligatoria.** Es la única columna sin la que una fila no se
puede cargar (`importar_visitas._parsear_fila`); el resto puede venir vacío --pasa
seguido con `Objetivo de compra`--, así que marcarlas todas como obligatorias, como
en canjes, describiría mal el archivo real.
"""
from dataclasses import dataclass
from io import BytesIO

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.services.estructura_archivo import ColumnaArchivo, EstructuraArchivo, GrupoColumnas
from app.services.importar_visitas import COLUMNAS_REQUERIDAS

HOJA = "VISITAS"

_AZUL = "3D3EA8"
_CORAL = "F4545A"


@dataclass(frozen=True)
class Columna:
    nombre: str
    grupo: str
    obligatoria: bool
    ayuda: str
    ancho: int = 22


# El orden es exactamente el de `COLUMNAS_REQUERIDAS`, que es el que la carga
# verifica y el que trae el archivo de la consola.
COLUMNAS: tuple[Columna, ...] = (
    Columna("Propiedad", "Propiedad", True, "Identifica la propiedad visitada. Es la única columna sin la que la fila no se carga.", 40),
    Columna("Dirección", "Propiedad", False, "Calle y comuna de la propiedad.", 30),
    Columna("Tipo", "Propiedad", False, "Texto libre tal como lo escribe la consola: Usada, Nueva.", 16),
    Columna("Mercado", "Propiedad", False, "Texto libre: Mercado secundario, Mercado primario.", 20),

    Columna("Cliente", "Cliente", False, "Quién pidió la visita.", 26),
    Columna("RUT", "Cliente", False, "Tal como viene de la consola, sin validar formato.", 16),
    Columna("Objetivo de compra", "Cliente", False, "Texto libre. Viene vacío en varias filas.", 24),

    Columna("Fecha/hora solicitada", "Solicitud", False, "Cuándo se agendó la visita.", 20),
    Columna("Etapa", "Solicitud", False, "Texto libre: Solicitada, y las que agregue la consola.", 16),
    Columna("Corredor", "Solicitud", False, "Quién de ViveProp o de la corredora aliada quedó a cargo.", 26),
    Columna("Solicitada el", "Solicitud", False, "Cuándo se generó la solicitud en la consola.", 20),
)

NOMBRES = tuple(c.nombre for c in COLUMNAS)


def estructura_importacion() -> EstructuraArchivo:
    """Las 11 columnas del export, agrupadas, para mostrarlas en pantalla."""
    grupos: list[GrupoColumnas] = []
    for col in COLUMNAS:
        if not grupos or grupos[-1].nombre != col.grupo:
            grupos.append(GrupoColumnas(nombre=col.grupo, columnas=[]))
        grupos[-1].columnas.append(
            ColumnaArchivo(nombre=col.nombre, obligatoria=col.obligatoria, ayuda=col.ayuda)
        )

    return EstructuraArchivo(
        titulo="Importar visitas",
        origen=(
            "El .xlsx que se exporta desde la consola de administración. No se llena "
            "a mano: la plantilla de acá sirve para comparar encabezados cuando la "
            "carga falla y no se entiende por qué."
        ),
        fila=(
            "Una fila es una solicitud de visita. El archivo no trae un identificador "
            "único, así que la carga arma uno propio con Propiedad + Cliente + RUT + "
            "Fecha/hora solicitada: si esos cuatro datos coinciden con una visita que ya "
            "está, se actualiza en vez de duplicarla."
        ),
        grupos=grupos,
        valores=[],
        notas=[
            "Los nombres de las columnas tienen que ser exactos. Si falta alguna, "
            "no se carga nada y el error dice cuál.",
            "Solo Propiedad no puede venir vacía. Las demás columnas sí, y quedan "
            "en blanco en la tabla.",
            "Reimportar el mismo archivo no duplica: actualiza Etapa y Corredor de "
            "las visitas que ya estaban, que son los datos que cambian con el tiempo.",
        ],
    )


def generar_plantilla() -> bytes:
    """El .xlsx vacío con los 11 encabezados exactos, en la fila 1.

    Sin fila de grupos: es donde `importar_visitas` los busca, igual que en
    canjes. Coral marca la única columna obligatoria, `Propiedad`; el resto va
    en azul.
    """
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = HOJA

    for i, col in enumerate(COLUMNAS, start=1):
        celda = hoja.cell(row=1, column=i, value=col.nombre)
        celda.font = Font(bold=True, color="FFFFFF", size=9)
        celda.fill = PatternFill("solid", fgColor=_CORAL if col.obligatoria else _AZUL)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        hoja.column_dimensions[get_column_letter(i)].width = col.ancho

    hoja.row_dimensions[1].height = 42
    hoja.freeze_panes = "A2"

    buffer = BytesIO()
    libro.save(buffer)
    return buffer.getvalue()
